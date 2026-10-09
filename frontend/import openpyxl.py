import openpyxl
import pandas as pd

wb = openpyxl.load_workbook(
    'Corporate Sessions - Completed.xlsx', data_only=True
)

sessions_records = []
pnl_records = []

for sheetname in wb.sheetnames:
  ws = wb[sheetname]
  current_month = None

  for r in range(1, ws.max_row + 1):
    c1 = ws.cell(row=r, column=1).value
    c2 = ws.cell(row=r, column=2).value
    c3 = ws.cell(row=r, column=3).value
    c4 = ws.cell(row=r, column=4).value
    c5 = ws.cell(row=r, column=5).value

    # Identify month section breaks in Col A
    if str(c1).strip() in [
        'January',
        'February',
        'March',
        'April',
        'May',
        'June',
        'July',
        'August',
        'September',
        'October',
        'November',
        'December',
    ]:
      current_month = str(c1).strip()
      continue

    # Extract session transaction rows
    if c1 not in [None, 'Date'] and c2 is not None:
      sessions_records.append({
          'Company': sheetname,
          'Month_Logged': current_month,
          'Date': c1,
          'Patient': c2,
          'Therapist': c3,
          'Status': str(c4).strip() if c4 else 'UNKNOWN',
          'Amount': c5,
      })

    # Extract Side-Table P&L (Columns I, J, K, L)
    p_month = ws.cell(row=r, column=9).value
    p_alloc = ws.cell(row=r, column=10).value
    p_payout = ws.cell(row=r, column=11).value
    p_profit = ws.cell(row=r, column=12).value

    if str(p_month).strip() in [
        'January',
        'February',
        'March',
        'April',
        'May',
        'June',
        'July',
        'August',
        'September',
        'October',
        'November',
        'December',
    ]:
      pnl_records.append({
          'Company': sheetname,
          'Month': str(p_month).strip(),
          'Allocated_Retainer': p_alloc,
          'Related_Payouts': p_payout,
          'Profit_Loss': p_profit,
      })

# Export to clean multi-tab workbook
with pd.ExcelWriter('Corporate_Cleaned_Model.xlsx', engine='openpyxl') as writer:
  pd.DataFrame(sessions_records).to_excel(
      writer, sheet_name='Corporate_Sessions', index=False
  )
  pd.DataFrame(pnl_records).to_excel(
      writer, sheet_name='Corporate_PnL', index=False
  )

print('Cleaned file generated: Corporate_Cleaned_Model.xlsx')